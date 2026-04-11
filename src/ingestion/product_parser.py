# src/ingestion/product_parser.py

import re
import pdfplumber
from openpyxl import load_workbook
from src.core.schema import create_layer_record
from src.utils.id_generator import generate_id


class ProductExtractor:

    # ============================
    # Public Methods
    # ============================

    def extract_from_pdf(self, file_path):
        return self._parse_pdf(file_path)

    def extract_from_excel(self, file_path):
        return self._parse_excel(file_path)

    # ============================
    # PDF Parsing
    # ============================

    def _parse_pdf(self, file_path):

        products = []

        try:
            with pdfplumber.open(file_path) as pdf:
                full_text = ""
                for page in pdf.pages:
                    full_text += page.extract_text() or ""

            if not full_text.strip():
                return []

            full_text = self._clean_text(full_text)

            product_data = {}

            patterns = {
                "product_name": r"Product Name:\s*(.*)",
                "manufacturer": r"Manufacturer:\s*(.*)",
                "model_number": r"Model Number:\s*(.*)",
                "applicable_standards": r"Applicable Standard[s]?:\s*(.*)",
                "fire_rating_hours": r"Fire Rating\s*(\d+)",
                "length_mm": r"Length\s*(\d+)",
                "width_mm": r"Width\s*(\d+)",
                "depth_mm": r"Depth\s*(\d+)",
                "compressive_strength_mpa": r"Compressive Strength\s*(\d+)",
                "warranty_years": r"Warranty Period\s*(\d+)"
            }

            for key, pattern in patterns.items():
                match = re.search(pattern, full_text, re.IGNORECASE)
                if match:
                    product_data[key] = match.group(1).strip()

            if product_data:
                record = create_layer_record(
                    record_id=generate_id(),
                    entity_type="Product",
                    layer="L2",
                    category="Technical",
                    properties=product_data
                )
                products.append(record)

            return products

        except Exception as e:
            print("PDF parsing error:", e)
            return []

    # ============================
    # Excel Parsing
    # ============================

    def _parse_excel(self, file_path):

        products = []

        try:
            wb = load_workbook(file_path)
            ws = wb.active

            headers = [cell.value for cell in ws[1]]

            for row in ws.iter_rows(min_row=2, values_only=True):

                properties = dict(zip(headers, row))

                record = create_layer_record(
                    record_id=generate_id(),
                    entity_type="Product",
                    layer="L2",
                    category="Technical",
                    properties=properties
                )

                products.append(record)

            return products

        except Exception as e:
            print("Excel parsing error:", e)
            return []

    # ============================
    # Cleaner
    # ============================

    def _clean_text(self, text):
        text = re.sub(r'\s+', ' ', text)
        text = text.replace("•", "")
        return text.strip()